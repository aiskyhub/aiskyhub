"""Mechanical regression fixtures; these do not claim to validate fictional art quality."""
from __future__ import annotations

import copy
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import comic_pipeline as cp
import comic_sources as cs
import comic_layout as cl


def fixture_script(source):
    script = cp.load_json(Path(cp.__file__).resolve().parents[1] / 'assets' / 'script-template.json')
    script.update(outline='测试夹具的全部事件顺序保留。', ending='以提供的最后一句结束。')
    script['style'].update(genre='测试', look='清晰线条', palette='灰色', selection_reason='自动化机械测试',
                           width=900, height=1200, font_size=28)
    script['characters'] = [{'id': 'char-a', 'name': '甲', 'aliases': [], 'importance': 'major',
        'narrative': {key: '测试人物' for key in ('goal', 'motivation', 'voice', 'arc')},
        'visual': {key: '测试设定' for key in ('face_shape', 'eyes', 'brows', 'nose_mouth', 'body', 'posture', 'hair', 'age')},
        'source_facts': [], 'design_notes': ['机械测试夹具，不代表实际人物图']}]
    script['settings'] = [{'id': 'room-a', 'description': '测试房间'}]
    body = [u for u in source['units'] if u['kind'] == 'body']
    for chapter in source['chapters']:
        if chapter['has_body']:
            script['scenes'].append({'id': 'scene-' + chapter['id'], 'chapter_id': chapter['id'], 'setting_id': 'room-a'})
    state = {'form': 'base', 'costume': 'coat-a', 'injuries': [], 'items': [], 'location': 'room-a', 'knowledge': []}
    for i, unit in enumerate(body, 1):
        event_id, panel_id = f'ev{i}', f'p{i}'
        script['events'].append({'id': event_id, 'description': unit['text'], 'source_unit_ids': [unit['id']]})
        script['panels'].append({'id': panel_id, 'chapter_id': unit['chapter_id'], 'scene_id': 'scene-' + unit['chapter_id'],
            'source_unit_ids': [unit['id']], 'event_ids': [event_id], 'cast': ['char-a'],
            'action': unit['text'], 'shot': '中景', 'space': '人物位于测试房间中', 'expression': '平静',
            'state_before': {'char-a': copy.deepcopy(state)}, 'state_after': {'char-a': copy.deepcopy(state)},
            'dialogue': [{'kind': 'caption', 'text': unit['text']}]})
        script['pages'].append({'id': f'page{i}', 'chapter_id': unit['chapter_id'], 'panel_ids': [panel_id], 'columns': 1})
    return script


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='comic-tests-')
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / '中文作品'
        self.source = self.base / '原文.txt'
        self.source.write_text('第1章 来信\n甲拿起信封。\n第2章 回信\n甲写下回答。\n', encoding='utf-8')
        self.invoke('init', source=[str(self.source)], title='通用测试')

    def invoke(self, name, **kwargs):
        from argparse import Namespace
        kwargs.setdefault('scope', None)
        return cp.run(Namespace(command=name, project=str(self.root), **kwargs))

    def json_file(self, value, stem='data'):
        import uuid
        path = self.base / (stem + '-' + uuid.uuid4().hex[:8] + '.json')
        cp.atomic_json(path, value)
        return str(path)

    def prepare_script(self):
        self.invoke('confirm-source', note='已检查夹具来源与边界')
        project = cp.project_load(self.root)
        for chapter in project['source']['chapters']:
            if chapter['has_body']:
                self.invoke('mark-read', chapter=chapter['id'], note='已读取夹具正文：' + chapter['title'])
        script = fixture_script(project['source'])
        self.invoke('set-script', file=self.json_file(script))
        return script

    def add_reviews(self):
        project = cp.project_load(self.root)
        for kind, checks in cp.REVIEW_CHECKS.items():
            report = {'script_hash': cp.digest(project['script']),
                'reviewed_chapter_ids': [c['id'] for c in project['source']['chapters'] if c['has_body']],
                'checks': {k: True for k in checks}, 'evidence': '自动化夹具报告，只测试结构与关卡。', 'issues': []}
            self.invoke('review', kind=kind, file=self.json_file(report))

    def locked(self):
        self.prepare_script()
        self.add_reviews()
        self.invoke('lock-script')

    def image_file(self, color='white'):
        from PIL import Image
        import uuid
        path = self.base / ('mechanical-' + uuid.uuid4().hex[:8] + '.png')
        Image.new('RGB', (600, 400), color).save(path)
        return str(path)

    def qa(self, keys):
        return self.json_file({'checks': {k: True for k in keys},
                               'evidence': '机械夹具，仅用于文件/流程回归，非真实图像验收。'})

    def reference(self):
        self.invoke('register-reference', characters=['char-a'], file=self.image_file(), qa=self.qa(cp.REFERENCE_CHECKS))

    def accept_all(self, colors=None):
        self.reference()
        for panel in cp.project_load(self.root)['script']['panels']:
            prompt = self.base / 'prompt.txt'
            prompt.write_text('机械测试，不调用图像服务', encoding='utf-8')
            attempt = self.invoke('begin-panel', panel=panel['id'], prompt=str(prompt))
            self.invoke('finish-panel', panel=panel['id'], attempt=attempt['attempt'],
                        file=self.image_file((colors or {}).get(panel['id'], '#b8c9d2')), qa=self.qa(cp.PANEL_CHECKS))

    def exported(self):
        self.locked()
        self.accept_all()
        self.invoke('compose', font=None)
        self.review_and_export()

    def review_and_export(self):
        layout = cp.project_load(self.root)['layout']
        report = {'input_hash': layout['input_hash'], 'reviewed_page_ids': [p['id'] for p in layout['pages']],
                  'checks': {k: True for k in cp.LAYOUT_CHECKS}, 'evidence': '机械页面夹具检查'}
        self.invoke('review-layout', file=self.json_file(report))
        self.invoke('export')

    def mixed_script(self, direction='ltr', max_height=6000):
        self.root = self.base / ('混合分格-' + direction)
        source = self.base / (direction + '.txt')
        source.write_text('第1章 场景\n甲进入房间。\n甲打开信封。\n甲读完信。\n甲走到窗前。\n', encoding='utf-8')
        self.invoke('init', source=[str(source)], title='混合分格机械测试')
        script = self.prepare_script()
        script['style'].update(format='strip', reading_direction=direction, max_segment_height=max_height)
        for panel in script['panels']:
            panel['dialogue'] = []
        script['pages'] = [{'id': 'page-mixed', 'chapter_id': script['panels'][0]['chapter_id'],
                            'panel_ids': ['p1', 'p2', 'p3', 'p4'], 'rows': [['p1'], ['p2', 'p3'], ['p4']]}]
        self.invoke('set-script', file=self.json_file(script))
        return script

    def accepted_mixed(self, direction='ltr', max_height=6000):
        self.mixed_script(direction, max_height)
        self.add_reviews()
        self.invoke('lock-script')
        self.accept_all({'p1': '#ca5362', 'p2': '#428d6b', 'p3': '#467cc2', 'p4': '#d3a348'})

    def color_bounds(self, image, color):
        from PIL import Image, ImageColor
        mask = Image.new('1', image.size)
        rgb = ImageColor.getrgb(color)
        pixels = image.load()
        mask.putdata([pixels[x, y] == rgb for y in range(image.height) for x in range(image.width)])
        bounds = mask.getbbox()
        self.assertIsNotNone(bounds, 'Rendered illustration color missing: ' + color)
        return bounds

    def test_page_rows_reject_missing_reordered_duplicate_or_malformed_panels(self):
        script = self.mixed_script()
        invalid = [[], [[]], None, ['p1', 'p2', 'p3', 'p4'],
                   [['p1', 'p2', 'p3'], ['p4']], [['p2', 'p1'], ['p3', 'p4']],
                   [['p1'], ['p2', 'p3']], [['p1'], ['p2', 'p2'], ['p4']],
                   [['p1'], ['p2', 3], ['p4']]]
        for rows in invalid:
            with self.subTest(rows=rows):
                candidate = copy.deepcopy(script)
                candidate['pages'][0]['rows'] = rows
                self.invoke('set-script', file=self.json_file(candidate))
                self.assertTrue(any('Page rows' in error for error in self.invoke('check-script')['errors']))

    def test_mixed_rows_render_full_width_pairs_and_both_reading_directions(self):
        from PIL import Image
        for direction in ('ltr', 'rtl'):
            with self.subTest(direction=direction):
                self.accepted_mixed(direction)
                self.invoke('compose', font=None)
                layout = cp.project_load(self.root)['layout']
                self.assertEqual(1, len(layout['pages']))
                self.assertEqual(['p1', 'p2', 'p3', 'p4'], layout['pages'][0]['panel_ids'])
                with Image.open(cp.inside(self.root, layout['pages'][0]['path'])) as image:
                    a, b, c, d = [self.color_bounds(image, color) for color in
                                   ('#ca5362', '#428d6b', '#467cc2', '#d3a348')]
                self.assertEqual((a[0], a[2]), (d[0], d[2]))
                self.assertEqual(b[2] - b[0], c[2] - c[0])
                self.assertGreater(a[2] - a[0], b[2] - b[0])
                self.assertEqual(b[1], c[1])
                self.assertLess(a[3], b[1])
                self.assertLess(b[3], d[1])
                self.assertEqual(b[0] < c[0], direction == 'ltr')
                for bounds in (a, b, c, d):
                    self.assertAlmostEqual((bounds[2] - bounds[0]) / (bounds[3] - bounds[1]), 1.5, delta=0.02)
                self.review_and_export()
                self.assertEqual(1, self.invoke('verify-export')['page_count'])

    def test_mixed_rows_split_at_row_boundaries_within_segment_limit(self):
        self.accepted_mixed(max_height=1000)
        self.invoke('compose', font=None)
        pages = cp.project_load(self.root)['layout']['pages']
        self.assertEqual([['p1', 'p2', 'p3'], ['p4']], [page['panel_ids'] for page in pages])
        self.assertTrue(all(page['height'] <= 1000 for page in pages))
        self.review_and_export()
        self.assertEqual(2, self.invoke('verify-export')['page_count'])

    def test_columns_fallback_renders_last_single_panel_at_full_width(self):
        from PIL import Image
        script = self.mixed_script()
        script['pages'] = [{'id': 'page-grid', 'chapter_id': script['panels'][0]['chapter_id'],
                            'panel_ids': ['p1', 'p2', 'p3'], 'columns': 2},
                           {'id': 'page-last', 'chapter_id': script['panels'][0]['chapter_id'], 'panel_ids': ['p4']}]
        self.invoke('set-script', file=self.json_file(script))
        self.add_reviews()
        self.invoke('lock-script')
        self.accept_all({'p1': '#ca5362', 'p2': '#428d6b', 'p3': '#467cc2', 'p4': '#d3a348'})
        self.invoke('compose', font=None)
        pages = cp.project_load(self.root)['layout']['pages']
        with Image.open(cp.inside(self.root, pages[0]['path'])) as image:
            a, b, c = [self.color_bounds(image, color) for color in ('#ca5362', '#428d6b', '#467cc2')]
        self.assertEqual(a[1], b[1])
        self.assertLess(a[3], c[1])
        self.assertGreater(c[2] - c[0], a[2] - a[0])

    def test_visual_plan_change_invalidates_only_affected_painting(self):
        self.locked()
        self.accept_all()
        script = cp.project_load(self.root)['script']
        script['panels'][0]['visual_plan'] = {'focal_point': '信封', 'depth': '前景手部，中景人物，简化背景'}
        self.invoke('set-script', file=self.json_file(script))
        self.add_reviews()
        self.invoke('lock-script')
        status = self.invoke('status')
        self.assertEqual(1, status['panels_accepted'])
        self.assertEqual('p1', status['next_panel'])

    def test_layout_version_invalidates_exports_and_page_cache_but_reuses_art(self):
        self.exported()
        before = cp.project_load(self.root)['layout']
        with patch.object(cl, 'LAYOUT_VERSION', cl.LAYOUT_VERSION + 1):
            with self.assertRaises(cp.GateError):
                self.invoke('verify-export')
            self.assertEqual(2, self.invoke('status')['panels_accepted'])
            self.invoke('compose', font=None)
            after = cp.project_load(self.root)['layout']
            self.assertNotEqual(before['input_hash'], after['input_hash'])
            self.assertNotEqual(before['pages'][0]['path'], after['pages'][0]['path'])
            self.review_and_export()
            self.assertEqual(2, self.invoke('verify-export')['page_count'])

    def test_pre_art_gate_rejects_empty_and_partial_script(self):
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')
        script = self.prepare_script()
        script['panels'].pop()
        script['pages'].pop()
        self.invoke('set-script', file=self.json_file(script))
        self.assertTrue(any('Unmapped' in x or 'Chapter has no' in x for x in self.invoke('check-script')['errors']))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')

    def test_three_current_reviews_required(self):
        self.prepare_script()
        project = cp.project_load(self.root)
        for kind in ('coverage', 'continuity'):
            report = {'script_hash': cp.digest(project['script']), 'reviewed_chapter_ids': ['ch000001', 'ch000002'],
                      'checks': {k: True for k in cp.REVIEW_CHECKS[kind]}, 'evidence': '机械审查测试'}
            self.invoke('review', kind=kind, file=self.json_file(report))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')
        self.add_reviews()
        self.invoke('lock-script')
        self.assertTrue(self.invoke('assert-art')['allowed'])
        self.assertTrue((self.root / 'full-script.md').is_file())

    def test_source_modified_invalidates_lock(self):
        self.locked()
        self.source.write_text(self.source.read_text(encoding='utf-8') + '新增事件。', encoding='utf-8')
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')
        self.assertFalse(self.invoke('status')['script_locked'])

    def test_index_tampering_rejected(self):
        self.locked()
        project = cp.project_load(self.root)
        project['source']['units'][1]['text'] = '偷偷删改正文'
        cp.save(self.root, project)
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')

    def test_script_change_invalidates_old_reviews(self):
        self.locked()
        script = cp.project_load(self.root)['script']
        script['panels'][0]['action'] += '（新镜头）'
        self.invoke('set-script', file=self.json_file(script))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')

    def test_duplicate_ids_and_wrong_speaker_rejected(self):
        script = self.prepare_script()
        script['panels'][1]['id'] = script['panels'][0]['id']
        script['panels'][0]['dialogue'] = [{'kind': 'speech', 'speaker': 'unknown', 'text': '你好'}]
        self.invoke('set-script', file=self.json_file(script))
        errors = self.invoke('check-script')['errors']
        self.assertTrue(any('duplicate IDs' in x for x in errors))
        self.assertTrue(any('speaker not in cast' in x for x in errors))

    def test_state_transition_needs_source_evidence(self):
        script = self.prepare_script()
        script['panels'][1]['state_before']['char-a']['costume'] = 'coat-b'
        script['panels'][1]['state_after']['char-a']['costume'] = 'coat-b'
        self.invoke('set-script', file=self.json_file(script))
        self.assertTrue(any('Unexplained' in x for x in self.invoke('check-script')['errors']))
        script['panels'][1]['state_transitions'] = [{'character_id': 'char-a', 'fields': ['costume'],
                                                  'reason': '夹具换装', 'source_unit_ids': script['panels'][1]['source_unit_ids']}]
        self.invoke('set-script', file=self.json_file(script))
        self.assertEqual([], self.invoke('check-script')['errors'])

    def test_major_unresolved_review_rejected(self):
        self.prepare_script()
        project = cp.project_load(self.root)
        report = {'script_hash': cp.digest(project['script']), 'reviewed_chapter_ids': ['ch000001', 'ch000002'],
                  'checks': {k: True for k in cp.REVIEW_CHECKS['coverage']}, 'evidence': '发现遗漏',
                  'issues': [{'severity': 'major', 'resolved': False}]}
        with self.assertRaises(cp.GateError):
            self.invoke('review', kind='coverage', file=self.json_file(report))

    def test_attempt_limit_and_pending_no_double_generation(self):
        self.locked()
        self.reference()
        prompt = self.base / 'prompt.txt'
        prompt.write_text('机械测试', encoding='utf-8')
        for i in range(3):
            attempt = self.invoke('begin-panel', panel='p1', prompt=str(prompt))
            with self.assertRaises(cp.GateError):
                self.invoke('begin-panel', panel='p1', prompt=str(prompt))
            self.invoke('fail-panel', panel='p1', attempt=attempt['attempt'], reason='机械失败测试')
            prompt.write_text('微小提示词变动' + str(i), encoding='utf-8')
        with self.assertRaises(cp.GateError):
            self.invoke('begin-panel', panel='p1', prompt=str(prompt))

    def test_reference_or_prompt_file_change_invalidates_art(self):
        self.locked()
        self.accept_all()
        self.assertEqual(2, self.invoke('status')['panels_accepted'])
        project = cp.project_load(self.root)
        path = cp.inside(self.root, project['art']['panels']['p1'][0]['prompt_path'])
        path.write_text('changed', encoding='utf-8')
        self.assertEqual(1, self.invoke('status')['panels_accepted'])
        reference = cp.inside(self.root, project['art']['references'][0]['path'])
        reference.write_bytes(b'changed')
        self.assertEqual(0, self.invoke('status')['panels_accepted'])

    def test_dialogue_change_reuses_paintings_but_requires_new_lock(self):
        self.locked()
        self.accept_all()
        script = cp.project_load(self.root)['script']
        script['panels'][0]['dialogue'][0]['text'] += '新增排版文字'
        self.invoke('set-script', file=self.json_file(script))
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')
        self.add_reviews()
        self.invoke('lock-script')
        self.assertEqual(2, self.invoke('status')['panels_accepted'])
        prompt = self.base / 'prompt.txt'
        prompt.write_text('此提示不会被调用', encoding='utf-8')
        self.assertTrue(self.invoke('begin-panel', panel='p1', prompt=str(prompt))['already_accepted'])

    def test_exports_real_order_hashes_completion_and_missing_file(self):
        self.exported()
        self.assertEqual(2, self.invoke('verify-export')['page_count'])
        self.assertFalse(self.invoke('status')['complete'])
        project = cp.project_load(self.root)
        report = {'input_hash': project['layout']['input_hash'],
                  'checks': {k: True for k in ['source_scope', 'story_complete', 'visual_consistency', 'exports_opened']},
                  'evidence': '机械夹具验证，非真实漫画视觉测试'}
        self.invoke('complete', file=self.json_file(report))
        self.assertTrue(self.invoke('status')['complete'])
        pdf = next(x for x in project['exports']['files'] if x['kind'] == 'pdf')
        cp.inside(self.root, pdf['path']).unlink()
        self.assertFalse(self.invoke('status')['complete'])

    def test_path_escape_rejected(self):
        with self.assertRaises(cp.GateError):
            cp.inside(self.root, '../outside.png')

    def test_unicode_encodings_duplicate_headings_and_empty_chapter(self):
        for encoding in ('utf-16', 'gb18030', 'utf-8-sig'):
            file = self.base / (encoding + '.txt')
            file.write_bytes('前言\n说明文字。\n第1章 同号\n正文。\n第1章 同号\n下一段。\n第2章 空章\n'.encode(encoding))
            source = cs.extract([file])
            self.assertEqual(4, len(source['chapters']))
            self.assertEqual(4, len({c['id'] for c in source['chapters']}))
            self.assertFalse(source['chapters'][-1]['has_body'])
            self.assertEqual(1, len(source['issues']))

    def test_docx_table_and_epub_spine_order(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('第1章 正文')
        doc.add_paragraph('文本。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表内书信。'
        path = self.base / '正文.docx'
        doc.save(path)
        self.assertIn('表内书信。', [u['text'] for u in cs.extract([path])['units']])
        path = self.base / '正文.epub'
        with zipfile.ZipFile(path, 'w') as book:
            book.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
            book.writestr('OPS/book.opf', '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
                '<item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/>'
                '</manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
            book.writestr('OPS/a.xhtml', '<html><body><p>后段。</p></body></html>')
            book.writestr('OPS/b.xhtml', '<html><body><p>先段。</p><script>隐藏内容</script></body></html>')
        self.assertEqual(['先段。', '后段。'], [u['text'] for u in cs.extract([path])['units']])

    def test_pdf_unreadable_page_flagged(self):
        from reportlab.pdfgen import canvas
        path = self.base / '空白页.pdf'
        pdf = canvas.Canvas(str(path))
        pdf.drawString(30, 700, 'Chapter 1')
        pdf.drawString(30, 680, 'Story text.')
        pdf.showPage()
        pdf.showPage()
        pdf.save()
        source = cs.extract([path])
        self.assertTrue(any(i.get('locator', {}).get('page') == 2 for i in source['issues']))

    def test_missing_font_glyph_stops_layout(self):
        self.locked()
        self.accept_all()
        project = cp.project_load(self.root)
        project['title'] = '𐐷'
        cp.save(self.root, project)
        with self.assertRaises(cp.GateError):
            self.invoke('compose', font=None)


if __name__ == '__main__':
    unittest.main()
